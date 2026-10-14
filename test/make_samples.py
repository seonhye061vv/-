import openpyxl
S=["지원서접수","서류전형","인적성전형","실무면접","임원면접","처우협의","평판조회","내부품의진행","최종합격발표"]
rows=[("손님지원본부","계약관리부","계약관리부 계약보전 경력사원 채용",1),("경영기획본부","재무기획부","재무기획부 세무담당 경력사원 채용",1),("CEO 임원","감사부","감사부 내부감사인 경력사원 채용",5),("계리리스크본부","리스크관리부","리스크관리부 경력사원 채용",7),("영업본부","하이브리드사업부","하이브리드사업부 총무 경력사원 채용",5),("영업본부","하이브리드사업부","하이브리드사업부 지점장 경력사원 채용",5),("손님지원본부","계약관리부","계약관리부 콜센터 경력사원 채용",0),("경영기획본부","재무기획부","재무기획부 재무관리 경력사원 채용",0),("CEO 임원","소비자보호부","소비자보호부 민원담당 경력사원 채용",0)]
wb=openpyxl.Workbook();ws=wb.active;ws.title="채용현황"
ws["A1"]="본부";ws["B1"]="부서";ws["C1"]="채용공고명";ws["D1"]="채용 프로세스";ws.merge_cells("D1:L1")
for i,s in enumerate(S): ws.cell(2,4+i,s)
for r,(h,d,t,c) in enumerate(rows):
  ws.cell(3+r,1,h);ws.cell(3+r,2,d);ws.cell(3+r,3,t)
  for i in range(9):
    v="완료" if i<c else "진행 중" if i==c else None
    if d=="감사부" and i==2: v="해당사항없음"
    if d=="리스크관리부" and i==6: v="해당사항 없음"
    ws.cell(3+r,4+i,v)
p=wb.create_sheet("프로세스")
for i,s in enumerate(S): p.cell(2,2+i,s); p.cell(3,2+i,"해당사항없음")
q=wb.create_sheet("권한부여대상자");q["C1"]="대상자";q.merge_cells("C1:F1")
for r,(h,d,cs) in enumerate([("계리리스크본부","계약관리부","YD"),("영업지원본부","재무기획부","ID"),("CEO 임원","감사부","HD"),("영업본부","하이브리드사업부","SD"),("CEO 임원","소비자보호관리책임자","WD")]):
  q.cell(2+r,1,h);q.cell(2+r,2,d)
  for j,c in enumerate(cs): q.cell(2+r,3+j,c)
import sys,os; out=sys.argv[1] if len(sys.argv)>1 else "."; wb.save(os.path.join(out,"sample_ok.xlsx"))
ws["E4"]="진행";ws.cell(11,1,"CEO 임원");ws.cell(11,2,"감사부");ws.cell(11,3,"감사부 내부감사인 경력사원 채용")
wb.save(os.path.join(out,"sample_bad.xlsx"))
# 행 순서만 바꾼 파일
import copy
wb2=openpyxl.load_workbook(os.path.join(out,"sample_ok.xlsx")); ws2=wb2["채용현황"]
rows=[[c.value for c in r] for r in ws2.iter_rows(min_row=3,max_row=11)]
for i,row in enumerate(reversed(rows)):
  for j,v in enumerate(row): ws2.cell(3+i,1+j,v)
wb2.save(os.path.join(out,"sample_moved.xlsx"))
ws2.cell(3,3,ws2.cell(3,3).value+" (수정)")
wb2.save(os.path.join(out,"sample_renamed.xlsx"))
# 대상자를 실명으로 적은 권한 시트 + 마지막 단계까지 끝난 공고 1건
wb3=openpyxl.load_workbook(os.path.join(out,"sample_ok.xlsx"))
q3=wb3["권한부여대상자"]
names={"Y":"윤하나","D":"클로드","I":"이서진","H":"한도윤","S":"서민재","W":"우지아"}
for row in q3.iter_rows(min_row=2,min_col=3,max_col=6):
  for c in row:
    if c.value in names: c.value=names[c.value]
w3=wb3["채용현황"]
for i in range(9): w3.cell(9,4+i,"완료")
wb3.save(os.path.join(out,"sample_names.xlsx"))
# 권한부여대상자 명단만 담은 파일 (이름·이메일 쌍, 시트 이름 자유). 예시 이름·주소만 사용
wb4=openpyxl.Workbook(); q4=wb4.active; q4.title="시트1"
q4.append(["본부","부서","권한부여대상자1","이메일주소","권한부여대상자2","이메일주소","권한부여대상자3","이메일주소"])
q4.append(["CEO 임원","감사부","가나다 부장","gana.da@example.com","라마바 상무","rama.ba@example.com","라마바 상무","rama.ba@example.com"])
q4.append(["경영기획본부","재무기획부","사아자 부장","saa.ja@example.com","라마바 상무","rama.ba@example.com"])
q4.append(["영업본부","하이브리드사업부","차카타 팀장","chaka@example","라마바 상무","rama.ba@example.com"])
wb4.save(os.path.join(out,"sample_perm.xlsx"))
